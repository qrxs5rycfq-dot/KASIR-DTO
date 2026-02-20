from datetime import datetime, timedelta
from io import BytesIO

from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user

from extensions import limiter, csrf
from models import (
    db, User, Category, MenuItem, Table, Order, OrderItem, Payment,
    Income, Setting, Cart, CartItem, Discount, PendingPrint, Notification,
    Branch, BranchMenuStock, CashierShift, Expense, ExternalOrder, City, Brand
)
from utils import (
    role_required, permission_required,
    get_user_branch_id, get_default_branch_id, branch_filter,
    utc_now, format_currency, get_setting, get_gateway_config
)

views_bp = Blueprint('views', __name__, url_prefix='')


# ============================================
# DASHBOARD
# ============================================

@views_bp.route('/dashboard')
@login_required
def dashboard():
    # Get statistics
    today = datetime.now().date()
    yesterday = today - timedelta(days=1)
    
    # Get only completed/paid orders today
    today_orders = branch_filter(Order.query, Order).filter(
        db.func.date(Order.created_at) == today
    ).all()
    
    # Calculate income from paid orders only
    paid_orders_today = [o for o in today_orders if o.payment and o.payment.status == 'paid']
    total_income_today = sum(o.total for o in paid_orders_today)
    total_orders_today = len(paid_orders_today)
    
    # Yesterday's data for real growth percentages
    yesterday_orders = branch_filter(Order.query, Order).filter(
        db.func.date(Order.created_at) == yesterday
    ).all()
    paid_orders_yesterday = [o for o in yesterday_orders if o.payment and o.payment.status == 'paid']
    total_income_yesterday = sum(o.total for o in paid_orders_yesterday)
    total_orders_yesterday = len(paid_orders_yesterday)
    
    # Calculate real growth percentages
    def calc_growth(current, previous):
        if previous == 0:
            return 100.0 if current > 0 else 0.0
        return round(((current - previous) / previous) * 100, 1)
    
    income_growth = calc_growth(total_income_today, total_income_yesterday)
    orders_growth = calc_growth(total_orders_today, total_orders_yesterday)
    
    today_avg = total_income_today // total_orders_today if total_orders_today > 0 else 0
    yesterday_avg = total_income_yesterday // total_orders_yesterday if total_orders_yesterday > 0 else 0
    avg_growth = calc_growth(today_avg, yesterday_avg)
    
    # Get popular items from actual order statistics (last 30 days)
    # Query to find most ordered items
    popular_query = db.session.query(
        OrderItem.menu_item_id,
        db.func.sum(OrderItem.quantity).label('total_ordered')
    ).join(Order).filter(
        Order.created_at >= datetime.now() - timedelta(days=30),
        Order.status.in_(['completed', 'processing'])
    ).group_by(OrderItem.menu_item_id).order_by(
        db.func.sum(OrderItem.quantity).desc()
    ).limit(6).all()
    
    # Get menu items for popular items
    popular_item_ids = [item[0] for item in popular_query if item[0]]
    if popular_item_ids:
        popular_items = MenuItem.query.filter(MenuItem.id.in_(popular_item_ids)).all()
        # Sort by order count
        item_order = {item_id: idx for idx, (item_id, _) in enumerate(popular_query)}
        popular_items.sort(key=lambda x: item_order.get(x.id, 999))
    else:
        # Fallback to marked popular items if no orders yet
        popular_items = MenuItem.query.filter_by(is_popular=True).limit(6).all()
    
    # Get recent orders
    recent_orders = branch_filter(Order.query, Order).order_by(Order.created_at.desc()).limit(10).all()
    
    # Get tables status - based on active orders (branch-filtered)
    tables = branch_filter(Table.query.filter_by(is_active=True), Table).all()
    
    # Calculate occupied tables from current active orders
    active_orders = branch_filter(Order.query, Order).filter(
        Order.status.in_(['pending', 'processing']),
        Order.table_id.isnot(None)
    ).all()
    occupied_table_ids = set(o.table_id for o in active_orders)
    
    # Update tables status dynamically
    for table in tables:
        if table.id in occupied_table_ids:
            table.status = 'occupied'
        else:
            table.status = 'available'
    
    # Low stock items (stock <= 10) - per-branch
    bid = get_user_branch_id()
    if bid:
        low_stock_items_raw = db.session.query(MenuItem, BranchMenuStock).join(
            BranchMenuStock, BranchMenuStock.menu_item_id == MenuItem.id
        ).filter(
            BranchMenuStock.branch_id == bid,
            BranchMenuStock.is_available == True,
            BranchMenuStock.stock <= 10
        ).order_by(BranchMenuStock.stock.asc()).all()
        # Attach stock to menu item objects for template compatibility
        low_stock_items = []
        for mi, bms in low_stock_items_raw:
            mi._branch_stock = bms.stock
            low_stock_items.append(mi)
    else:
        low_stock_items = MenuItem.query.filter(
            MenuItem.is_available == True,
            MenuItem.stock <= 10
        ).order_by(MenuItem.stock.asc()).all()
        for mi in low_stock_items:
            mi._branch_stock = mi.stock
    
    # Per-branch breakdown for admin dashboard
    branch_stats = []
    is_admin = current_user.has_role('admin')
    if is_admin:
        all_branches = Branch.query.order_by(Branch.name).all()
        for branch in all_branches:
            b_orders = [o for o in Order.query.filter(
                Order.branch_id == branch.id,
                db.func.date(Order.created_at) == today
            ).all() if o.payment and o.payment.status == 'paid']
            branch_stats.append({
                'name': branch.name,
                'income': sum(o.total for o in b_orders),
                'orders': len(b_orders)
            })
    
    return render_template('dashboard.html',
                         total_income_today=total_income_today,
                         total_orders_today=total_orders_today,
                         income_growth=income_growth,
                         orders_growth=orders_growth,
                         avg_growth=avg_growth,
                         popular_items=popular_items,
                         recent_orders=recent_orders,
                         tables=tables,
                         low_stock_items=low_stock_items,
                         branch_stats=branch_stats,
                         is_admin=is_admin,
                         now=datetime.now())


# ============================================
# POS (Kasir)
# ============================================

@views_bp.route('/pos')
@login_required
@role_required('manager', 'kasir')
def pos():
    categories = branch_filter(Category.query.filter_by(is_active=True), Category).order_by(Category.order).all()
    menu_items = branch_filter(MenuItem.query.filter_by(is_available=True), MenuItem).all()
    tables = branch_filter(Table.query.filter_by(is_active=True), Table).all()
    
    return render_template('pos.html',
                         categories=categories,
                         menu_items=menu_items,
                         tables=tables,
                         now=datetime.now())


# ============================================
# Online Order (via QR code)
# ============================================

@views_bp.route('/order/online/<table_number>')
def online_order(table_number):
    table = Table.query.filter_by(number=table_number).first()
    if not table:
        flash('Meja tidak ditemukan!', 'danger')
        return redirect(url_for('auth.login'))
    
    categories = Category.query.filter_by(is_active=True).order_by(Category.order).all()
    menu_items = MenuItem.query.filter_by(is_available=True).all()
    
    return render_template('online_order.html',
                         table=table,
                         categories=categories,
                         menu_items=menu_items,
                         now=datetime.now())


# ============================================
# PROFILE
# ============================================

@views_bp.route('/profile')
@login_required
def profile():
    return render_template('profile.html', user=current_user)

@views_bp.route('/profile/update', methods=['POST'])
@login_required
def update_profile():
    full_name = request.form.get('full_name')
    email = request.form.get('email')
    phone = request.form.get('phone')
    
    # Check email uniqueness
    if email != current_user.email:
        if User.query.filter_by(email=email).first():
            flash('Email sudah digunakan!', 'danger')
            return redirect(url_for('views.profile'))
    
    current_user.full_name = full_name
    current_user.email = email
    current_user.phone = phone
    
    db.session.commit()
    flash('Profil berhasil diperbarui!', 'success')
    return redirect(url_for('views.profile'))

@views_bp.route('/profile/change-password', methods=['POST'])
@login_required
def profile_change_password():
    current_password = request.form.get('current_password')
    new_password = request.form.get('new_password')
    confirm_password = request.form.get('confirm_password')
    
    if not current_user.check_password(current_password):
        flash('Password saat ini salah!', 'danger')
        return redirect(url_for('views.profile'))
    
    if new_password != confirm_password:
        flash('Password baru tidak cocok!', 'danger')
        return redirect(url_for('views.profile'))
    
    if len(new_password) < 6:
        flash('Password minimal 6 karakter!', 'danger')
        return redirect(url_for('views.profile'))
    
    current_user.set_password(new_password)
    db.session.commit()
    flash('Password berhasil diubah!', 'success')
    return redirect(url_for('views.profile'))


# ============================================
# PRINTER STATION
# ============================================

@views_bp.route('/printer-station')
@login_required
def printer_station():
    """Dedicated printer station page - keep this open for reliable printing"""
    return render_template('printer_station.html')


# ============================================
# EXPENSES
# ============================================

@views_bp.route('/expenses')
@login_required
@role_required('admin', 'manager', 'kasir')
def expenses():
    """Expense tracking page"""
    from datetime import date
    
    # Get filter parameters
    month = request.args.get('month', date.today().strftime('%Y-%m'))
    category_filter = request.args.get('category', 'all')
    
    try:
        filter_year, filter_month = map(int, month.split('-'))
    except (ValueError, AttributeError):
        filter_year, filter_month = date.today().year, date.today().month
    
    # Build query
    query = branch_filter(Expense.query, Expense).filter(
        db.extract('year', Expense.date) == filter_year,
        db.extract('month', Expense.date) == filter_month
    )
    
    if category_filter != 'all':
        query = query.filter(Expense.category == category_filter)
    
    expense_list = query.order_by(Expense.date.desc(), Expense.created_at.desc()).all()
    
    # Calculate totals
    total_expenses = sum(e.amount for e in expense_list)
    category_totals = {}
    for e in expense_list:
        label = Expense.CATEGORIES.get(e.category, e.category)
        category_totals[label] = category_totals.get(label, 0) + e.amount
    
    return render_template('expenses.html',
                         expenses=expense_list,
                         total_expenses=total_expenses,
                         category_totals=category_totals,
                         expense_categories=Expense.CATEGORIES,
                         current_month=month,
                         category_filter=category_filter,
                         active_page='expenses',
                         now=datetime.now())

@views_bp.route('/expenses/add', methods=['POST'])
@login_required
@role_required('admin', 'manager', 'kasir')
def expense_add():
    """Add a new expense"""
    from datetime import date as date_type
    
    expense_date = request.form.get('date')
    category = request.form.get('category', '').strip()
    description = request.form.get('description', '').strip()
    amount = request.form.get('amount', '0')
    notes = request.form.get('notes', '').strip()
    
    if not description or not category:
        flash('Deskripsi dan kategori wajib diisi.', 'danger')
        return redirect(url_for('views.expenses'))
    
    try:
        amount = int(amount)
        if amount <= 0:
            raise ValueError
    except (ValueError, TypeError):
        flash('Jumlah harus angka positif.', 'danger')
        return redirect(url_for('views.expenses'))
    
    try:
        parsed_date = date_type.fromisoformat(expense_date) if expense_date else date_type.today()
    except ValueError:
        parsed_date = date_type.today()
    
    expense = Expense(
        date=parsed_date,
        category=category,
        description=description,
        amount=amount,
        notes=notes,
        user_id=current_user.id,
        branch_id=get_default_branch_id()
    )
    db.session.add(expense)
    db.session.commit()
    
    flash(f'Pengeluaran "{description}" berhasil ditambahkan!', 'success')
    return redirect(url_for('views.expenses'))

@views_bp.route('/expenses/<int:expense_id>/delete', methods=['POST'])
@login_required
@role_required('admin', 'manager')
def expense_delete(expense_id):
    """Delete an expense"""
    expense = db.session.get(Expense, expense_id)
    if not expense:
        flash('Pengeluaran tidak ditemukan.', 'danger')
        return redirect(url_for('views.expenses'))
    
    db.session.delete(expense)
    db.session.commit()
    
    flash('Pengeluaran berhasil dihapus.', 'success')
    return redirect(url_for('views.expenses'))


# ============================================
# KITCHEN
# ============================================

@views_bp.route('/kitchen')
@login_required
@role_required('manager', 'koki', 'kasir')
def kitchen():
    """Kitchen display page for branch staff"""
    return render_template('kitchen.html')


# ============================================
# REPORTS & ANALYTICS
# ============================================

@views_bp.route('/reports')
@login_required
@role_required('admin', 'manager', 'kasir')
def reports():
    return render_template('reports.html')

@views_bp.route('/analytics')
@login_required
@role_required('admin', 'manager', 'kasir')
def analytics():
    """Comprehensive analytics dashboard with real data and percentages.
    Supports filtering by city_id, brand_id, and branch_id query parameters."""
    today = datetime.now().date()
    yesterday = today - timedelta(days=1)
    week_ago = today - timedelta(days=7)
    two_weeks_ago = today - timedelta(days=14)
    month_start = today.replace(day=1)
    last_month_start = (month_start - timedelta(days=1)).replace(day=1)
    last_month_end = month_start - timedelta(days=1)
    
    # ── Filter parameters (for owner/admin: can filter by city/brand/branch) ──
    filter_city_id = request.args.get('city_id', '', type=str).strip()
    filter_brand_id = request.args.get('brand_id', '', type=str).strip()
    filter_branch_id = request.args.get('branch_id', '', type=str).strip()
    
    # Query Pusat branch once for legacy NULL branch_id handling
    pusat_branch = Branch.query.filter_by(code='PUSAT').first()
    pusat_branch_id = pusat_branch.id if pusat_branch else None
    
    def _includes_pusat(branch_ids):
        """Check if branch_ids list includes Pusat (for legacy NULL data handling)."""
        return pusat_branch_id is not None and pusat_branch_id in branch_ids
    
    def analytics_filter(query):
        """Apply branch/city/brand filter for analytics. Regular users see their branch only."""
        bid = get_user_branch_id()
        if bid is not None:
            # Non-owner: locked to their branch
            return query.filter(Order.branch_id == bid)
        # Owner: apply optional filters
        if filter_branch_id and filter_branch_id.isdigit():
            target_id = int(filter_branch_id)
            # Also include legacy orders with NULL branch_id if filtering for Pusat
            if target_id == pusat_branch_id:
                return query.filter(db.or_(Order.branch_id == target_id, Order.branch_id.is_(None)))
            return query.filter(Order.branch_id == target_id)
        if filter_brand_id and filter_brand_id.isdigit():
            branch_ids = [id for (id,) in Branch.query.with_entities(Branch.id).filter_by(brand_id=int(filter_brand_id)).all()]
            if not branch_ids:
                return query.filter(Order.id == None)
            if _includes_pusat(branch_ids):
                return query.filter(db.or_(Order.branch_id.in_(branch_ids), Order.branch_id.is_(None)))
            return query.filter(Order.branch_id.in_(branch_ids))
        if filter_city_id and filter_city_id.isdigit():
            branch_ids = [id for (id,) in Branch.query.with_entities(Branch.id).filter_by(city_id=int(filter_city_id)).all()]
            if not branch_ids:
                return query.filter(Order.id == None)
            if _includes_pusat(branch_ids):
                return query.filter(db.or_(Order.branch_id.in_(branch_ids), Order.branch_id.is_(None)))
            return query.filter(Order.branch_id.in_(branch_ids))
        return query  # Owner with no filter = all data
    
    def analytics_expense_filter(query):
        """Apply branch/city/brand filter for expenses."""
        bid = get_user_branch_id()
        if bid is not None:
            return query.filter(Expense.branch_id == bid)
        if filter_branch_id and filter_branch_id.isdigit():
            target_id = int(filter_branch_id)
            if target_id == pusat_branch_id:
                return query.filter(db.or_(Expense.branch_id == target_id, Expense.branch_id.is_(None)))
            return query.filter(Expense.branch_id == target_id)
        if filter_brand_id and filter_brand_id.isdigit():
            branch_ids = [id for (id,) in Branch.query.with_entities(Branch.id).filter_by(brand_id=int(filter_brand_id)).all()]
            if not branch_ids:
                return query.filter(Expense.id == None)
            if _includes_pusat(branch_ids):
                return query.filter(db.or_(Expense.branch_id.in_(branch_ids), Expense.branch_id.is_(None)))
            return query.filter(Expense.branch_id.in_(branch_ids))
        if filter_city_id and filter_city_id.isdigit():
            branch_ids = [id for (id,) in Branch.query.with_entities(Branch.id).filter_by(city_id=int(filter_city_id)).all()]
            if not branch_ids:
                return query.filter(Expense.id == None)
            if _includes_pusat(branch_ids):
                return query.filter(db.or_(Expense.branch_id.in_(branch_ids), Expense.branch_id.is_(None)))
            return query.filter(Expense.branch_id.in_(branch_ids))
        return query
    
    # Helper: calculate growth percentage
    def growth_pct(current, previous):
        if previous == 0:
            return 100.0 if current > 0 else 0.0
        return round(((current - previous) / previous) * 100, 1)
    
    # ── Today vs Yesterday ──
    today_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) == today
    ).all()
    yesterday_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) == yesterday
    ).all()
    
    today_paid = [o for o in today_orders if o.payment and o.payment.status == 'paid']
    yesterday_paid = [o for o in yesterday_orders if o.payment and o.payment.status == 'paid']
    
    today_revenue = sum(o.total for o in today_paid)
    yesterday_revenue = sum(o.total for o in yesterday_paid)
    revenue_growth = growth_pct(today_revenue, yesterday_revenue)
    
    today_order_count = len(today_paid)
    yesterday_order_count = len(yesterday_paid)
    order_growth = growth_pct(today_order_count, yesterday_order_count)
    
    today_avg = today_revenue // today_order_count if today_order_count > 0 else 0
    yesterday_avg = yesterday_revenue // yesterday_order_count if yesterday_order_count > 0 else 0
    avg_growth = growth_pct(today_avg, yesterday_avg)
    
    today_cancelled = len([o for o in today_orders if o.status == 'cancelled'])
    today_cancel_rate = round(today_cancelled / len(today_orders) * 100, 1) if today_orders else 0
    
    # ── This Week vs Last Week ──
    this_week_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) >= week_ago,
        db.func.date(Order.created_at) <= today
    ).all()
    last_week_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) >= two_weeks_ago,
        db.func.date(Order.created_at) < week_ago
    ).all()
    
    this_week_paid = [o for o in this_week_orders if o.payment and o.payment.status == 'paid']
    last_week_paid = [o for o in last_week_orders if o.payment and o.payment.status == 'paid']
    
    week_revenue = sum(o.total for o in this_week_paid)
    last_week_revenue = sum(o.total for o in last_week_paid)
    week_growth = growth_pct(week_revenue, last_week_revenue)
    
    # ── This Month vs Last Month ──
    this_month_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) >= month_start,
        db.func.date(Order.created_at) <= today
    ).all()
    last_month_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) >= last_month_start,
        db.func.date(Order.created_at) <= last_month_end
    ).all()
    
    this_month_paid = [o for o in this_month_orders if o.payment and o.payment.status == 'paid']
    last_month_paid = [o for o in last_month_orders if o.payment and o.payment.status == 'paid']
    
    month_revenue = sum(o.total for o in this_month_paid)
    last_month_revenue = sum(o.total for o in last_month_paid)
    month_growth = growth_pct(month_revenue, last_month_revenue)
    
    # ── Payment Method Breakdown (this month) ──
    payment_methods = {}
    for o in this_month_paid:
        method = o.payment.payment_method if o.payment else 'unknown'
        if method not in payment_methods:
            payment_methods[method] = {'count': 0, 'total': 0}
        payment_methods[method]['count'] += 1
        payment_methods[method]['total'] += o.total
    
    total_payment_count = sum(v['count'] for v in payment_methods.values())
    for method in payment_methods:
        payment_methods[method]['percentage'] = round(
            (payment_methods[method]['count'] / total_payment_count * 100), 1
        ) if total_payment_count > 0 else 0
    
    # ── Order Source Breakdown (this month) ──
    source_breakdown = {}
    for o in this_month_paid:
        src = o.source or 'pos'
        if src not in source_breakdown:
            source_breakdown[src] = {'count': 0, 'total': 0}
        source_breakdown[src]['count'] += 1
        source_breakdown[src]['total'] += o.total
    
    total_source_count = sum(v['count'] for v in source_breakdown.values())
    for src in source_breakdown:
        source_breakdown[src]['percentage'] = round(
            (source_breakdown[src]['count'] / total_source_count * 100), 1
        ) if total_source_count > 0 else 0
    
    # ── Order Type Breakdown (dine_in / takeaway / online) ──
    type_breakdown = {}
    for o in this_month_paid:
        otype = o.order_type or 'dine_in'
        if otype not in type_breakdown:
            type_breakdown[otype] = {'count': 0, 'total': 0}
        type_breakdown[otype]['count'] += 1
        type_breakdown[otype]['total'] += o.total
    
    total_type_count = sum(v['count'] for v in type_breakdown.values())
    for otype in type_breakdown:
        type_breakdown[otype]['percentage'] = round(
            type_breakdown[otype]['count'] / total_type_count * 100, 1
        ) if total_type_count > 0 else 0
    
    # ── Category Performance (this month) ──
    category_perf = {}
    for o in this_month_paid:
        for item in o.items:
            cat_name = item.menu_item.category.name if item.menu_item and item.menu_item.category else 'Tanpa Kategori'
            if cat_name not in category_perf:
                category_perf[cat_name] = {'count': 0, 'qty': 0, 'total': 0}
            category_perf[cat_name]['count'] += 1
            category_perf[cat_name]['qty'] += item.quantity
            category_perf[cat_name]['total'] += item.subtotal
    
    total_cat_revenue = sum(v['total'] for v in category_perf.values())
    for cat in category_perf:
        category_perf[cat]['percentage'] = round(
            (category_perf[cat]['total'] / total_cat_revenue * 100), 1
        ) if total_cat_revenue > 0 else 0
    
    # Sort by revenue desc
    category_perf = dict(sorted(category_perf.items(), key=lambda x: x[1]['total'], reverse=True))
    
    # ── Top 10 Menu Items (this month) ──
    top_items = {}
    for o in this_month_paid:
        for item in o.items:
            name = item.name
            if name not in top_items:
                top_items[name] = {'qty': 0, 'total': 0}
            top_items[name]['qty'] += item.quantity
            top_items[name]['total'] += item.subtotal
    
    top_items = dict(sorted(top_items.items(), key=lambda x: x[1]['qty'], reverse=True)[:10])
    total_item_qty = sum(v['qty'] for v in top_items.values())
    for name in top_items:
        top_items[name]['percentage'] = round(
            (top_items[name]['qty'] / total_item_qty * 100), 1
        ) if total_item_qty > 0 else 0
    
    # ── Hourly Order Distribution (this week) ──
    hourly_dist = {h: 0 for h in range(24)}
    for o in this_week_paid:
        if o.created_at:
            hourly_dist[o.created_at.hour] += 1
    
    peak_hour = max(hourly_dist, key=hourly_dist.get) if any(hourly_dist.values()) else 12
    
    # ── Daily Revenue Trend (last 30 days) ──
    thirty_days_ago = today - timedelta(days=30)
    month_all_orders = analytics_filter(Order.query).filter(
        db.func.date(Order.created_at) >= thirty_days_ago,
        db.func.date(Order.created_at) <= today
    ).all()
    
    daily_trend = {}
    for o in month_all_orders:
        if o.payment and o.payment.status == 'paid':
            day_str = o.created_at.strftime('%Y-%m-%d')
            if day_str not in daily_trend:
                daily_trend[day_str] = {'revenue': 0, 'orders': 0}
            daily_trend[day_str]['revenue'] += o.total
            daily_trend[day_str]['orders'] += 1
    
    # Fill missing days with 0
    daily_labels = []
    daily_revenues = []
    daily_orders_list = []
    current_day = thirty_days_ago
    while current_day <= today:
        day_str = current_day.strftime('%Y-%m-%d')
        daily_labels.append(current_day.strftime('%d/%m'))
        daily_revenues.append(daily_trend.get(day_str, {}).get('revenue', 0))
        daily_orders_list.append(daily_trend.get(day_str, {}).get('orders', 0))
        current_day += timedelta(days=1)
    
    # ── Expenses this month ──
    expense_query = analytics_expense_filter(Expense.query.filter(
        db.func.date(Expense.date) >= month_start,
        db.func.date(Expense.date) <= today
    ))
    total_expenses = sum(e.amount for e in expense_query.all())
    net_profit = month_revenue - total_expenses
    profit_margin = round((net_profit / month_revenue * 100), 1) if month_revenue > 0 else 0
    
    return render_template('analytics.html',
        today_revenue=today_revenue,
        yesterday_revenue=yesterday_revenue,
        revenue_growth=revenue_growth,
        today_order_count=today_order_count,
        order_growth=order_growth,
        today_avg=today_avg,
        avg_growth=avg_growth,
        today_cancel_rate=today_cancel_rate,
        week_revenue=week_revenue,
        week_growth=week_growth,
        month_revenue=month_revenue,
        month_growth=month_growth,
        month_orders=len(this_month_paid),
        payment_methods=payment_methods,
        source_breakdown=source_breakdown,
        type_breakdown=type_breakdown,
        category_perf=category_perf,
        top_items=top_items,
        hourly_dist=hourly_dist,
        peak_hour=peak_hour,
        daily_labels=daily_labels,
        daily_revenues=daily_revenues,
        daily_orders_list=daily_orders_list,
        total_expenses=total_expenses,
        net_profit=net_profit,
        profit_margin=profit_margin,
        filter_city_id=filter_city_id,
        filter_brand_id=filter_brand_id,
        filter_branch_id=filter_branch_id,
        cities=City.query.order_by(City.name).all(),
        brands=Brand.query.order_by(Brand.name).all(),
        branches=Branch.query.filter_by(is_active=True).order_by(Branch.name).all()
    )


# ============================================
# INCOME REPORT
# ============================================

@views_bp.route('/reports/income')
@login_required
@role_required('admin', 'manager', 'kasir')
def income_report():
    # Get date range from query params
    start_date_str = request.args.get('start_date', datetime.now().strftime('%Y-%m-01'))
    end_date_str = request.args.get('end_date', datetime.now().strftime('%Y-%m-%d'))
    
    # Parse dates properly
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d') + timedelta(days=1) - timedelta(seconds=1)
    
    orders = branch_filter(Order.query, Order).filter(
        Order.created_at >= start_date,
        Order.created_at <= end_date
    ).all()
    
    paid_orders = [o for o in orders if o.payment and o.payment.status == 'paid']
    
    total_income = sum(o.total for o in paid_orders)
    total_orders = len(paid_orders)
    
    # Group by date
    daily_income = {}
    for order in paid_orders:
        date_str = order.created_at.strftime('%Y-%m-%d')
        if date_str not in daily_income:
            daily_income[date_str] = {'income': 0, 'orders': 0}
        daily_income[date_str]['income'] += order.total
        daily_income[date_str]['orders'] += 1
    
    return render_template('reports/income.html',
                         orders=paid_orders,
                         total_income=total_income,
                         total_orders=total_orders,
                         daily_income=daily_income,
                         start_date=start_date_str,
                         end_date=end_date_str)


# ============================================
# EXPORT PDF
# ============================================

@views_bp.route('/reports/export/pdf')
@login_required
@role_required('admin', 'manager', 'kasir')
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm, cm
    from reportlab.platypus import SimpleDocTemplate, Table as PDFTable, TableStyle, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    
    start_date = request.args.get('start_date', datetime.now().strftime('%Y-%m-01'))
    end_date = request.args.get('end_date', datetime.now().strftime('%Y-%m-%d'))
    
    orders = branch_filter(Order.query, Order).filter(
        Order.created_at >= start_date,
        Order.created_at <= end_date + ' 23:59:59'
    ).all()
    
    paid_orders = [o for o in orders if o.payment and o.payment.status == 'paid']
    total_revenue = sum(o.total for o in paid_orders)
    total_orders = len(paid_orders)
    avg_order = round(total_revenue / total_orders) if total_orders > 0 else 0
    
    # Payment method breakdown
    payment_methods = {}
    for o in paid_orders:
        method = o.payment.payment_method if o.payment else 'cash'
        payment_methods[method] = payment_methods.get(method, 0) + 1
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                           topMargin=1.5*cm, bottomMargin=1.5*cm,
                           leftMargin=1.5*cm, rightMargin=1.5*cm)
    elements = []
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'],
                                fontSize=20, textColor=colors.HexColor('#1a1a2e'),
                                spaceAfter=4)
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'],
                                   fontSize=10, textColor=colors.HexColor('#666666'))
    header_style = ParagraphStyle('SectionHeader', parent=styles['Heading2'],
                                 fontSize=13, textColor=colors.HexColor('#1a1a2e'),
                                 spaceBefore=14, spaceAfter=8)
    
    brand_color = colors.HexColor('#f97316')
    dark_bg = colors.HexColor('#1a1a2e')
    light_bg = colors.HexColor('#f8f9fa')
    
    # Header
    elements.append(Paragraph("Dapoer Teras Obor", title_style))
    elements.append(Paragraph(f"Laporan Penjualan | Periode: {start_date} s/d {end_date}", subtitle_style))
    branch_name = current_user.branch.name if current_user.branch_id and current_user.branch else "Semua Cabang"
    elements.append(Paragraph(f"Cabang: {branch_name} | Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
    elements.append(Spacer(1, 6))
    elements.append(HRFlowable(width="100%", thickness=2, color=brand_color))
    elements.append(Spacer(1, 10))
    
    # Summary cards as a table
    def fmt_rp(val):
        return f"Rp {val:,.0f}".replace(',', '.')
    
    summary_data = [
        ['Total Pendapatan', 'Jumlah Transaksi', 'Rata-rata Transaksi'],
        [fmt_rp(total_revenue), str(total_orders), fmt_rp(avg_order)]
    ]
    summary_table = PDFTable(summary_data, colWidths=[6*cm, 5.5*cm, 5.5*cm])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), dark_bg),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BACKGROUND', (0, 1), (-1, 1), light_bg),
        ('TEXTCOLOR', (0, 1), (-1, 1), dark_bg),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 14),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#dee2e6')),
        ('LINEBELOW', (0, 0), (-1, 0), 1, brand_color),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 12))
    
    # Payment methods breakdown
    if payment_methods:
        elements.append(Paragraph("Metode Pembayaran", header_style))
        pm_data = [['Metode', 'Jumlah', 'Persentase']]
        for method, count in sorted(payment_methods.items(), key=lambda x: -x[1]):
            pct = (count / total_orders * 100) if total_orders > 0 else 0
            pm_data.append([method.upper(), str(count), f"{pct:.1f}%"])
        pm_table = PDFTable(pm_data, colWidths=[7*cm, 5*cm, 5*cm])
        pm_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), dark_bg),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dee2e6')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light_bg]),
        ]))
        elements.append(pm_table)
        elements.append(Spacer(1, 12))
    
    # Transaction detail table
    elements.append(Paragraph("Detail Transaksi", header_style))
    data = [['No', 'Tanggal', 'Order ID', 'Customer', 'Metode', 'Total']]
    for i, order in enumerate(paid_orders, 1):
        data.append([
            str(i),
            order.created_at.strftime('%d/%m/%Y %H:%M'),
            order.order_number,
            (order.customer_name or 'Guest')[:20],
            (order.payment.payment_method if order.payment else '-').upper(),
            fmt_rp(order.total)
        ])
    
    # Total row
    data.append(['', '', '', '', 'TOTAL', fmt_rp(total_revenue)])
    
    col_widths = [1.2*cm, 3.2*cm, 3*cm, 3.5*cm, 2.5*cm, 3.6*cm]
    table = PDFTable(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), dark_bg),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        # Body
        ('FONTNAME', (0, 1), (-1, -2), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -2), 8),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, light_bg]),
        # Total row
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e8f5e9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, -1), (-1, -1), 9),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, dark_bg),
        # Alignment
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (-2, 0), (-2, -1), 'CENTER'),
        # Grid
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dee2e6')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#adb5bd')),
    ]))
    elements.append(table)
    
    # Footer
    elements.append(Spacer(1, 16))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#dee2e6')))
    footer_style = ParagraphStyle('Footer', parent=styles['Normal'],
                                  fontSize=8, textColor=colors.HexColor('#999999'),
                                  alignment=TA_CENTER)
    elements.append(Spacer(1, 4))
    elements.append(Paragraph("Dokumen ini digenerate otomatis oleh Sistem Kasir Dapoer Teras Obor", footer_style))
    
    doc.build(elements)
    buffer.seek(0)
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'laporan_{start_date}_{end_date}.pdf',
        mimetype='application/pdf'
    )


# ============================================
# EXPORT EXCEL
# ============================================

@views_bp.route('/reports/export/excel')
@login_required
@role_required('admin', 'manager', 'kasir')
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
    from openpyxl.utils import get_column_letter
    
    start_date = request.args.get('start_date', datetime.now().strftime('%Y-%m-01'))
    end_date = request.args.get('end_date', datetime.now().strftime('%Y-%m-%d'))
    
    orders = branch_filter(Order.query, Order).filter(
        Order.created_at >= start_date,
        Order.created_at <= end_date + ' 23:59:59'
    ).all()
    
    paid_orders = [o for o in orders if o.payment and o.payment.status == 'paid']
    total_revenue = sum(o.total for o in paid_orders)
    total_orders_count = len(paid_orders)
    avg_order = total_revenue // total_orders_count if total_orders_count > 0 else 0
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Penjualan"
    
    # Colors & styles
    brand_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
    dark_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    light_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    summary_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    dark_font = Font(color="1A1A2E", bold=True)
    header_font = Font(color="1A1A2E", bold=True, size=18)
    sub_font = Font(color="666666", size=10)
    thin_border = Border(
        left=Side(style='thin', color='DEE2E6'),
        right=Side(style='thin', color='DEE2E6'),
        top=Side(style='thin', color='DEE2E6'),
        bottom=Side(style='thin', color='DEE2E6')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    # Column widths
    col_widths = {'A': 6, 'B': 18, 'C': 16, 'D': 20, 'E': 14, 'F': 18, 'G': 14}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Header row 1 - Brand
    ws.merge_cells('A1:G1')
    cell = ws['A1']
    cell.value = 'DAPOER TERAS OBOR'
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    
    # Header row 2 - Report info
    ws.merge_cells('A2:G2')
    branch_name = current_user.branch.name if current_user.branch_id and current_user.branch else "Semua Cabang"
    ws['A2'] = f'Laporan Penjualan | {start_date} s/d {end_date} | Cabang: {branch_name}'
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='center')
    
    # Brand accent line
    for col in range(1, 8):
        cell = ws.cell(row=3, column=col)
        cell.fill = brand_fill
    ws.row_dimensions[3].height = 4
    
    # Summary section (row 5-6)
    summary_labels = ['Total Pendapatan', 'Jumlah Transaksi', 'Rata-rata Transaksi']
    summary_values = [total_revenue, total_orders_count, avg_order]
    summary_cols = [(1, 2), (3, 4), (5, 6)]
    
    for i, ((c1, c2), label, val) in enumerate(zip(summary_cols, summary_labels, summary_values)):
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        cell = ws.cell(row=5, column=c1, value=label)
        cell.fill = dark_fill
        cell.font = Font(color="FFFFFF", size=9)
        cell.alignment = center_align
        ws.cell(row=5, column=c2).fill = dark_fill
        
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
        val_cell = ws.cell(row=6, column=c1, value=val)
        val_cell.fill = summary_fill
        val_cell.font = Font(color="1A1A2E", bold=True, size=13)
        val_cell.alignment = center_align
        ws.cell(row=6, column=c2).fill = summary_fill
        if i != 1:  # Format as currency except count
            val_cell.number_format = '#,##0'
    
    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 28
    
    # Blank row
    # Detail header (row 8)
    ws.merge_cells('A8:G8')
    ws['A8'] = 'Detail Transaksi'
    ws['A8'].font = Font(color="1A1A2E", bold=True, size=12)
    
    # Column headers (row 9)
    headers = ['No', 'Tanggal', 'Order ID', 'Customer', 'Metode', 'Total', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=9, column=col, value=header)
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[9].height = 22
    
    # Data rows
    for row_idx, order in enumerate(paid_orders, 10):
        row_num = row_idx - 9
        bg = light_fill if row_num % 2 == 0 else PatternFill()
        
        cells_data = [
            row_num,
            order.created_at.strftime('%d/%m/%Y %H:%M'),
            order.order_number,
            order.customer_name or 'Guest',
            (order.payment.payment_method if order.payment else '-').upper(),
            order.total,
            'LUNAS'
        ]
        
        for col, value in enumerate(cells_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            if row_num % 2 == 0:
                cell.fill = bg
            if col == 1:
                cell.alignment = center_align
            elif col == 5 or col == 7:
                cell.alignment = center_align
            elif col == 6:
                cell.alignment = right_align
                cell.number_format = '#,##0'
    
    # Total row
    total_row = len(paid_orders) + 10
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    total_label = ws.cell(row=total_row, column=1, value='TOTAL')
    total_label.font = Font(bold=True, size=11)
    total_label.alignment = Alignment(horizontal='right', vertical='center')
    total_label.fill = green_fill
    total_label.border = thin_border
    for c in range(2, 6):
        ws.cell(row=total_row, column=c).fill = green_fill
        ws.cell(row=total_row, column=c).border = thin_border
    
    total_val = ws.cell(row=total_row, column=6, value=total_revenue)
    total_val.font = Font(bold=True, size=11)
    total_val.alignment = right_align
    total_val.fill = green_fill
    total_val.number_format = '#,##0'
    total_val.border = thin_border
    ws.cell(row=total_row, column=7).fill = green_fill
    ws.cell(row=total_row, column=7).border = thin_border
    ws.row_dimensions[total_row].height = 24
    
    # Footer
    footer_row = total_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=7)
    ws.cell(row=footer_row, column=1, 
            value=f'Digenerate otomatis oleh Sistem Kasir Dapoer Teras Obor | {datetime.now().strftime("%d/%m/%Y %H:%M")}')
    ws.cell(row=footer_row, column=1).font = Font(color="999999", size=8, italic=True)
    ws.cell(row=footer_row, column=1).alignment = Alignment(horizontal='center')
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'laporan_{start_date}_{end_date}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================
# PAYMENT PAGE
# ============================================

@views_bp.route('/payment/<int:order_id>')
@login_required
def payment_page(order_id):
    from flask import current_app

    order = db.session.get(Order, order_id)
    if not order:
        flash('Pesanan tidak ditemukan', 'error')
        return redirect(url_for('views.orders'))
    
    # Determine active gateway
    active_gateway = get_setting('active_payment_gateway', 'midtrans')
    
    # Only generate Midtrans snap token when midtrans is the active gateway
    snap_token = None
    if active_gateway == 'midtrans':
        if order.payment and order.payment.snap_token:
            snap_token = order.payment.snap_token
        elif order.payment and order.payment.status == 'pending':
            try:
                import midtransclient
                
                is_prod_val = get_gateway_config('midtrans_is_production', 'MIDTRANS_IS_PRODUCTION')
                snap = midtransclient.Snap(
                    is_production=is_prod_val.lower() == 'true' if is_prod_val else current_app.config.get('MIDTRANS_IS_PRODUCTION', False),
                    server_key=get_gateway_config('midtrans_server_key', 'MIDTRANS_SERVER_KEY') or current_app.config.get('MIDTRANS_SERVER_KEY', ''),
                    client_key=get_gateway_config('midtrans_client_key', 'MIDTRANS_CLIENT_KEY') or current_app.config.get('MIDTRANS_CLIENT_KEY', '')
                )
                
                param = {
                    "transaction_details": {
                        "order_id": f"DTO-{order.id}-{int(datetime.now().timestamp())}",
                        "gross_amount": int(order.total)
                    },
                    "customer_details": {
                        "first_name": current_user.full_name or current_user.username,
                        "email": current_user.email or f"{current_user.username}@dapoerterasobor.com"
                    },
                    "item_details": [{
                        "id": str(item.menu_item_id),
                        "price": int(item.price),
                        "quantity": item.quantity,
                        "name": item.menu_item.name[:50]
                    } for item in order.items]
                }
                
                transaction = snap.create_transaction(param)
                snap_token = transaction.get('token')
                
                order.payment.snap_token = snap_token
                db.session.commit()
            except Exception as e:
                print(f"Error generating snap token: {e}")
    
    # Get Midtrans client key for Snap.js (DB-first, config fallback)
    midtrans_client_key = get_gateway_config('midtrans_client_key', 'MIDTRANS_CLIENT_KEY') or current_app.config.get('MIDTRANS_CLIENT_KEY', '')
    midtrans_is_prod_val = get_gateway_config('midtrans_is_production', 'MIDTRANS_IS_PRODUCTION')
    midtrans_is_production = midtrans_is_prod_val.lower() == 'true' if midtrans_is_prod_val else current_app.config.get('MIDTRANS_IS_PRODUCTION', False)
    
    return render_template('payment.html', 
                         order=order, 
                         snap_token=snap_token or '',
                         auto_pay=bool(snap_token),
                         active_gateway=active_gateway,
                         midtrans_client_key=midtrans_client_key,
                         midtrans_is_production=midtrans_is_production,
                         config=current_app.config)


# ============================================
# ORDERS
# ============================================

@views_bp.route('/orders')
@login_required
@role_required('admin', 'manager', 'kasir')
def orders():
    status_filter = request.args.get('status', 'all')
    
    query = branch_filter(Order.query, Order).order_by(Order.created_at.desc())
    
    if status_filter != 'all':
        query = query.filter_by(status=status_filter)
    
    orders = query.limit(100).all()
    
    return render_template('orders.html', orders=orders, status_filter=status_filter)


# ============================================
# ERROR HANDLERS
# ============================================

@views_bp.app_errorhandler(404)
def page_not_found(e):
    return render_template('errors/404.html'), 404

@views_bp.app_errorhandler(500)
def internal_server_error(e):
    return render_template('errors/500.html'), 500
